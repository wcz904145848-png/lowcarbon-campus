# -*- coding: utf-8 -*-
"""生成论文 docx：碳种校园"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FIG = r"D:\HuaHua\u\低碳游戏化程序\Collab\_probe\paper"
OUT = r"D:\HuaHua\u\低碳游戏化程序\碳种校园_论文.docx"

HEI, SONG = "黑体", "宋体"


def set_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for k in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rFonts.set(qn(k), name)


def para(doc, text, font=SONG, size=12, bold=False, align=None,
         indent=True, before=0, after=6, line=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), font, size, bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.4
    pf.keep_with_next = True
    pf.keep_together = True
    set_font(p.add_run(text), HEI, 14 if level == 1 else 12)
    return p


def _no_borders(table):
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblPrChange")


def fig_row(doc, items, width=Cm(4.6)):
    i = 0
    while i < len(items):
        pair = items[i:i + 2]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        for j, (path, _cap) in enumerate(pair):
            if j:
                p.add_run("　　")
            p.add_run().add_picture(path, width=width)
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        set_font(c.add_run("　　　".join(cap for _, cap in pair)), SONG, 9)
        i += 2

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = sec.bottom_margin = Cm(2.5)

# ===== 标题 =====
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
set_font(p.add_run("碳种校园：校园低碳行为科普的游戏化信息可视化设计策略研究"), HEI, 20)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("RESEARCH ON GAMIFIED INFORMATION VISUALIZATION DESIGN STRATEGY "
                   "FOR CAMPUS LOW-CARBON BEHAVIOR POPULARIZATION"), SONG, 9)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
set_font(p.add_run("（作者姓名）　（学校·学院）"), SONG, 10.5)

# ===== 摘要 =====
ABS = ("摘要：在“双碳”目标与绿色校园建设的背景下，高校学生普遍认同低碳理念，却存在“知道但不做、"
       "做了不知道省了多少”的知行断层。本文以校园低碳行为的科普信息可视化为研究对象，提出游戏化机制"
       "与信息可视化相融合的设计路径，并以移动端应用《碳种校园》为实践载体进行验证。研究首先梳理现有"
       "校园低碳科普在信息组织、行为引导与反馈机制三方面的不足；继而从视觉信息的同步处理优势、图优"
       "效应与沉浸体验三个层面，论证游戏化语境下信息可视化的作用机理；据此提出四条设计策略：以“价值"
       "可验算”替代口号式宣传、以空间锚定替代列表罗列、以生长过程替代数值进度、以“只读账本”隔离"
       "认知与娱乐。设计实践将 19 类校园低碳行为抽象为“活动量 ×（基准因子 − 实际因子）”的减排"
       "计算模型，构建校园漫游与农田养成两段式场景，以可行走的 2.5D 校园地图承载任务情境，以作物"
       "生长映射行为累积。结果表明，该设计在不牺牲科学严谨性的前提下，把抽象的碳减排量转译为可感知、"
       "可积累、可回溯的视觉经验，为校园低碳科普提供了一种可复用的设计范式。")
para(doc, ABS, SONG, 9, indent=True, line=1.4, after=4)
para(doc, "关键词：信息可视化；游戏化设计；低碳行为；校园科普；行为闭环；2.5D 插画",
     SONG, 9, indent=True, line=1.4, after=4)
para(doc, "中图分类号：J52　　　文献标识码：A", SONG, 9, indent=False, line=1.4, after=10)

ABS_EN = ("Abstract: Under the goals of carbon peaking and carbon neutrality and the construction of "
          "green campuses, university students generally accept low-carbon concepts, yet a gap persists "
          "between knowing and doing: many are willing but do not act, and those who act cannot tell how "
          "much they have saved. Taking the popular-science visualization of campus low-carbon behavior as "
          "its object, this paper proposes a design path that integrates gamification mechanisms with "
          "information visualization, and validates it through a mobile application, Carbon Seed Campus. "
          "The study first reviews three deficiencies of existing campus low-carbon popularization in "
          "information organization, behavioral guidance and feedback; it then argues for the mechanism of "
          "visualization within a gamified context from three aspects: the synchronous processing advantage "
          "of visual information, the picture-superiority effect, and immersive experience. On this basis, "
          "four design strategies are proposed: replacing slogan-style propaganda with verifiable value, "
          "replacing list-style enumeration with spatial anchoring, replacing numeric progress with growth "
          "process, and separating cognition from entertainment through a read-only ledger. The practice "
          "abstracts nineteen categories of campus low-carbon behavior into a calculation model of quantity "
          "multiplied by the difference between baseline and actual emission factors, and constructs a "
          "two-stage scene of campus roaming and farmland cultivation.")
para(doc, ABS_EN, SONG, 8.5, indent=True, line=1.3, after=4)
para(doc, "Keywords: Information visualization; Gamification design; Low-carbon behavior; Campus "
          "popular science; Behavioral loop; 2.5D illustration", SONG, 8.5, indent=True, line=1.3, after=12)

# ===== 引言 =====
heading(doc, "引言", 1)
for t in [
    "校园是城市低碳转型的微观实验场。教室照明、通勤方式、餐饮结构与一次性用品消耗，构成了学生日常"
    "碳排放的主要来源，也构成了最容易干预的行为场景。近年来，“绿色校园”“零碳校园”建设在高校持续"
    "推进，宣传栏、主题班会、志愿服务等科普形式不断丰富，学生对低碳理念的认同度显著提升。",
    "然而认同并不等于行动。笔者在前期调研中发现两类典型现象：其一，学生能够复述“节约用电”“绿色"
    "出行”等口号，却难以判断自己的某一具体行为究竟减少了多少排放，因而缺乏持续行动的内在理由；其二，"
    "即便参与者完成了行为，也往往止于一次性的打卡，缺少将零散行为累积为可见成果的机制。换言之，"
    "现有科普完成了“知识传递”，却未完成“价值显影”与“行为固化”。",
    "信息可视化设计在此具有不可替代的作用。它把抽象的量、关系与过程转译为图形，使受众在阅读时能够"
    "同步处理多条信息，从而降低理解门槛。但仅凭静态信息图，仍难以解决“看懂了却不做”的问题。游戏化"
    "所提供的目标、规则、反馈与循环，恰好补足了从认知到行为的最后一环。",
    "本文的工作是：将游戏化机制与信息可视化方法结合，提出针对校园低碳科普的四条设计策略，并以自行"
    "设计开发的竖屏移动端作品《碳种校园》加以实现与检验，讨论该路径在科学严谨性与游戏趣味性之间取得"
    "平衡的具体做法。",
]:
    para(doc, t)

# ===== 一 =====
heading(doc, "一、校园低碳科普的现状与问题", 1)
heading(doc, "（一）校园低碳行为的实施与科普现状", 2)
for t in [
    "校园低碳行为大致可分为四类：出行替代（步行、骑行替代机动车）、能源节约（随手关灯、空调调温）、"
    "物资循环（旧书、旧衣、闲置物品流转）与消费替代（自带水杯、减少一次性餐具、植物性饮食）。这四类"
    "行为在校园中发生频率高、单次减排量小，属于典型的“高频微行为”。",
    "微行为的特点决定了科普的难点。单次行为的减排量多在 0.1—2 kg CO₂e 量级，远低于公众的感知阈值；"
    "而行为本身又高度依赖场景触发——学生并非不知道“骑车更环保”，而是在赶时间去教学楼的具体情境中"
    "没有被提醒。因此，校园低碳科普的关键不在于“告知更多知识”，而在于“在正确的时刻，以可感知的"
    "尺度，呈现行为的价值”。",
]:
    para(doc, t)

heading(doc, "（二）现有可视化设计存在的三类问题", 2)
for t in [
    "对现有校园低碳科普物料（宣传栏、折页、公众号长图、打卡小程序）的梳理显示，三类问题较为普遍。",
    "其一，信息组织上的“知识罗列化”。多数物料以“低碳生活十条”“节能减排小妙招”的形式平铺条目，"
    "条目之间缺乏逻辑关系与优先级，受众读完仍不知道“对我而言，哪一件事最值得做”。图表中常见的是"
    "装饰性的树叶、地球剪影，而非承载数据的图形编码，视觉元素与信息量之间不成比例。",
    "其二，行为引导上的“场景缺席”。科普内容脱离具体的时间与空间，学生很难在真实情境中调用。反观"
    "优秀的科普可视化实践，往往以场景叙事组织信息，将抽象的处置流程还原为可想象的操作动作，从而在"
    "受众头脑中预演行为。",
    "其三，反馈机制上的“成果不可见”。打卡类产品通常只记录次数，不呈现累积效果；学生完成 10 次绿色"
    "出行，界面上仍是“10 次”，而非“累计减少 2.0 kg 排放”。缺乏可累积、可回望的成果形态，行为难以"
    "沉淀为习惯。",
]:
    para(doc, t)

# ===== 二 =====
heading(doc, "二、游戏化与信息可视化融合的优势", 1)
heading(doc, "（一）游戏化机制：把“知道”变成“做到”", 2)
for t in [
    "游戏化的核心并非添加积分与徽章，而是移植游戏得以成立的结构：明确的目标、清晰的规则、即时的反馈"
    "与自愿参与的循环。当行为被置于这样的结构中，动机从外部规范转向内部成就，重复行为的发生成本随之"
    "降低。",
    "对于高频微行为，游戏化提供的是“累积的容器”。单次减排 0.2 kg 几乎不构成意义，但 30 次累积为 "
    "6 kg，并转化为一株作物的成熟，行为便获得了叙事上的重量。这种“以小博大”的价值转换，正是微行为"
    "激励的关键。",
]:
    para(doc, t)

heading(doc, "（二）可视化在游戏化语境中的三重优势", 2)
for t in [
    "在游戏化框架内，信息可视化承担的不再是“说明”，而是“驱动”。其优势体现在三个层面。",
    "第一，视觉信息的同步处理优势。受众解读视觉信息时能够对多个维度同时加工，而文字必须线性推进。"
    "当“减排量、折算成长值、账本余额”三者以同一屏内的图形并置呈现时，用户无需计算即可建立因果关联，"
    "认知负荷显著降低。",
    "第二，图形信息的记忆优势。知觉的“图优效应”使图形化信息在延迟回忆中保持更高的留存率。将“步行"
    "替代汽车”抽象为一条由脚印延伸为路径的图形，比文字描述更容易在下次出行决策时被唤起。",
    "第三，可视化的沉浸体验。2.5D 轴测视角介于图示与写实之间，既保留场景的可辨识性，又具备图示的"
    "概括性，能够在不牺牲信息清晰度的前提下营造空间沉浸感，使受众从“旁观者”转为“在场者”。",
]:
    para(doc, t)

# ===== 三 =====
heading(doc, "三、校园低碳科普信息可视化的设计策略", 1)
para(doc, "基于上述分析，本文提出四条相互支撑的设计策略。它们共同回答一个问题：如何让一次微小的低碳"
          "行为，在视觉上获得与它科学价值相称的重量。")

heading(doc, "（一）以“价值可验算”替代口号式宣传", 2)
for t in [
    "低碳科普最容易滑向口号。要建立可信度，必须让每一个数字都能被用户自己复算。具体做法是把减排量"
    "表达为一个开放的算式：活动量 ×（基准情景因子 − 实际情景因子），并在界面上完整展示基准值、实际值"
    "与代入过程。用户看到的不是“本次减排 0.20 kg”这样一个结论，而是一条可以质疑、可以核对的推导链。",
    "这一策略的代价是必须承认不确定性。作品在每个算式下方标注“基准系数为演示假设，非正式校园核算"
    "系数”，把科普的边界如实交代。诚实地暴露假设，反而增强了整体的可信度。",
]:
    para(doc, t)

heading(doc, "（二）以空间锚定替代列表罗列", 2)
for t in [
    "把低碳任务放进一张可以行走的校园地图，是本设计在信息组织上的核心选择。任务不再是列表中的一行，"
    "而是依附于具体建筑的空间事件：图书馆对应无纸化课堂，食堂对应自带水杯与植物性饮食，宿舍对应随手"
    "关灯与旧衣回收。",
    "空间锚定的价值在于两点。其一，它提供了情境线索，用户在真实校园中经过该建筑时，地图上的标记会成为"
    "行为提醒；其二，它把“科普阅读”转化为“场景探索”，用户为获得信息而移动，移动本身构成注意力的"
    "投入，信息的接收从被动转为主动。",
]:
    para(doc, t)
fig_row(doc, [(os.path.join(FIG, "fig_map.jpg"), "图 1　校园地图漫游场景"),
              (os.path.join(FIG, "fig_tasks.jpg"), "图 2　低碳任务列表")])

heading(doc, "（三）以生长过程替代数值进度", 2)
for t in [
    "数值进度条传达的是“还差多少”，生长过程传达的是“已经变成什么”。在《碳种校园》中，行为的累积"
    "不被表示为百分比，而被表示为作物的五个形态阶段：种子、幼苗、小树、成树、盛放。每一次有效的低碳"
    "行为都换算为一次灌溉或施肥，推进形态跃迁。",
    "这一转换带来了两点收益。其一，多阶段的形态比连续进度条更易被记忆与描述，用户会以“我的樱花开了”"
    "而非“我的进度到了 80%”来叙述自己的成果；其二，形态本身构成了视觉奖励，成熟时的花冠与结果，"
    "是任何数字都无法提供的即时满足。",
]:
    para(doc, t)
fig_row(doc, [(os.path.join(FIG, "fig_calc.jpg"), "图 3　减排计算与双账户反馈"),
              (os.path.join(FIG, "fig_farm.jpg"), "图 4　农田养成场景")])

heading(doc, "（四）以“只读账本”隔离认知与娱乐", 2)
for t in [
    "游戏化最大的风险是让严肃数据被娱乐逻辑侵蚀：一旦减排量可以兑换、可以消耗，它的科学性便不复存在。"
    "为此，本设计设置了两个物理隔离的账户：碳账本（碳减排记录）与游戏资源（金币、成长值、作物）。",
    "碳账本只增不减，任何游戏内消费都不会触碰它；游戏资源可以自由流通，但不反过来写入碳账本。"
    "界面在计算页并列展示两者的走向——“获得成长值 +2”与“碳账本（只读）→ 0.20 kg”，并在说明文字中"
    "明确写出“碳账本不会因游戏行为被消耗”。这一设计使科普的严肃性与游戏的自由度得以并存。",
]:
    para(doc, t)
fig_row(doc, [(os.path.join(FIG, "fig_bag.jpg"), "图 5　收获物进入背包并出售")], width=Cm(4.9))

# ===== 四 =====
heading(doc, "四、设计实践：《碳种校园》", 1)
heading(doc, "（一）信息抽取与数值建模", 2)
for t in [
    "设计的第一步是把低碳知识转译为可计算的字段，而非可排版的文本。研究梳理出 19 类校园低碳行为，"
    "为每一类建立两条因子：基准情景（如乘汽车 0.200 kg/km）与实际情景（如步行 0.000 kg/km），"
    "两者之差乘以活动量即得减排量。因子表覆盖出行、用电、物资与餐饮四个维度，所有数值集中存放于单一"
    "数据文件中，便于按学校实际情况替换。",
    "与减排量并行的是游戏侧的换算规则：每 0.1 kg 减排折算 1 点成长值，成长值推动等级提升，等级每升"
    "一级解锁两块新的农田；金币作为独立的经济资源，用于购买种子、浇水与施肥。两套数值通过“一次行为"
    "同时产出两种结果”的规则耦合，但互不侵入。",
]:
    para(doc, t)

heading(doc, "（二）视觉风格定位", 2)
for t in [
    "整体视觉语言定为 2.5D 轴测插画：场景以等距视角铺陈，建筑与植被保留手绘质感，光照从上左方向下"
    "右方推进，形成统一的体积感。选择轴测而非透视，是因为轴测不产生近大远小的强变形，同一屏内不同"
    "位置的信息具备可比性，符合可视化对“可度量”的要求。",
    "色彩上采用低饱和的校园绿为基底，以米白承载信息卡片，用暖黄（金币）与深绿（成长）区分两类反馈。"
    "角色以四方向行走帧呈现——正面、背面、左侧、右侧各八帧，在移动时按运动主轴切换，兼顾表现力与"
    "资源体积。",
]:
    para(doc, t)

heading(doc, "（三）交互结构", 2)
for t in [
    "作品采用竖屏 1080×1920 的设计基准，按屏幕尺寸等比缩放，包含校园与农田两个主场景与若干功能页。",
    "校园场景中，用户点击地面即可行走，系统在可通行区域上以 A* 算法规划路径并逐点绘制引导标记；"
    "走近建筑时标签高亮，轻点即可进入对应的低碳任务。任务页以卡片列表呈现，每张卡片同时给出行为描述、"
    "估算减排量与金币奖励（图 2）。进入某一任务后，计算页将算式、基准值、实际值与双账户走向一次性"
    "呈现（图 3）。",
    "农田场景以全屏等距底图承载 16 块田地。点空地直接弹出种子选择，点已种植的作物则在植株上方浮出"
    "信息气泡，只显示名称、阶段与进度，不承载操作；浇水、施肥、收获、背包等操作统一收敛到屏幕底部的"
    "图标操作栏。这一分工解决了移动端“气泡遮挡植株”与“手指遮挡操作区”的两难。收获物先进入背包，"
    "出售后才转化为金币，形成“收获—持有—变现”的完整经济闭环（图 5）。",
]:
    para(doc, t)

heading(doc, "（四）实现与验证", 2)
for t in [
    "作品以纯前端技术实现：HTML、CSS 与原生 JavaScript，不依赖任何框架与构建工具。校园场景使用 "
    "Canvas 逐帧绘制，碰撞检测由一张灰度掩码图逐像素采样完成，寻路在 4 像素网格上运行 A*；界面层"
    "使用 DOM，以便文本清晰与可访问性。",
    "为适配移动端播放环境，静态资源经三步优化：剔除未被引用的素材、将大幅底图转为 WebP、合并四方向"
    "行走帧为单张精灵表。部署包体积由 19.7 MB 降至 3.2 MB，在 4G 网络下的首屏加载时间约 2 秒。"
    "作品同时提供网页版与 Android 安装包两种形态，后者以内置 WebView 加载本地资源，可完全离线运行。",
    "在可用性自测中，重点验证了三项：其一，点击寻路在草地、道路、广场等不同地形上的成功率，通过掩码"
    "自检脚本完成四方向连通性测试；其二，田地块位与底图泥土格的像素级对齐，通过截图色彩检测反复校正；"
    "其三，存档在版本变更后的容错，通过版本号比对丢弃不兼容存档而非报错。",
]:
    para(doc, t)

# ===== 结论 =====
heading(doc, "结论", 1)
for t in [
    "本文围绕“如何让校园低碳科普从知识传递走向行为固化”这一问题，提出并实践了一条游戏化与信息可视化"
    "融合的设计路径。研究认为，面向高频微行为的科普可视化，其难点不在信息量，而在价值的可感知性："
    "必须让用户看见自己的行为值多少、累积成什么、将来会怎样。",
    "据此形成的四条策略——价值可验算、空间锚定、生长叙事、只读账本——分别回应了可信度、情境性、"
    "累积感与严肃性四个维度。在《碳种校园》的实践中，这些策略被落实为减排算式、可行走的校园地图、"
    "五阶段作物形态与双账户结构，验证了其可行性与自洽性。",
    "研究仍存在局限。其一，减排因子取自公开资料，尚未与本校准数体系对接，数值仅具演示意义；其二，"
    "目前仅完成单机形态，缺少真实用户的行为数据回流，策略的有效性尚待规模化的对照实验检验。后续工作"
    "将围绕因子本地化、多人协作场景与行为数据的纵向追踪展开。",
]:
    para(doc, t)

# ===== 参考文献 =====
heading(doc, "参考文献", 1)
REFS = [
    "[1] 樊荣, 陈紫琪, 陈鹏. 社区垃圾分类的科普信息可视化设计策略研究[J]. 设计, 2024, 37(12): 140-144.",
    "[2] 中共中央, 国务院. 关于完整准确全面贯彻新发展理念做好碳达峰碳中和工作的意见[Z]. 2021-09-22.",
    "[3] 李拓. 信息可视化：让公众“看得见”疫情的设计方法[J]. 装饰, 2020(2): 38-45.",
    "[4] 周华清, 李小霞. 中外医学类期刊科学可视化设计比较研究[J]. 出版科学, 2021, 29(4): 32-41.",
    "[5] 易晓丽, 王瑾. 冬季运动项目可视化设计探索——以高山滑雪运动项目为例[J]. 设计, 2019, 32(5): 33-35.",
    "[6] 杨梅, 李航. 格式塔心理学视知觉原理在扁平化设计中的应用[J]. 包装工程, 2019, 40(8): 72-75.",
    "[7] 樊荣, 李笑晗, 陈鹏. 后疫情时代下新冠疫苗信息可视化设计研究[J]. 设计, 2022, 35(21): 64-68.",
    "[8] 简·麦戈尼格尔. 游戏改变世界[M]. 闾佳, 译. 杭州: 浙江人民出版社, 2012.",
    "[9] 米哈里·契克森米哈赖. 心流：最优体验心理学[M]. 张定绮, 译. 北京: 中信出版社, 2017.",
    "[10] SCHELL J. The Art of Game Design: A Book of Lenses[M]. Boca Raton: CRC Press, 2008.",
    "[11] TUFTE E R. The Visual Display of Quantitative Information[M]. 2nd ed. Cheshire: Graphics Press, 2001.",
    "[12] 朱红宇. 青少年科普信息可视化设计研究[D]. 无锡: 江南大学, 2022.",
]
for r in REFS:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(3)
    set_font(p.add_run(r), SONG, 9)

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT) // 1024, "KB")
